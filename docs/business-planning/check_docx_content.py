from pathlib import Path
from docx import Document

path = Path(r"D:\NewTrevl\docs\TravelOS业务功能规划建议.docx")
doc = Document(path)
texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
table_texts = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
full_text = "\n".join(texts + table_texts)
missing_features = [f"F{i:02d}" for i in range(1, 15) if f"F{i:02d}" not in full_text]
placeholders = [token for token in ("TODO", "[Company Name]", "待补充", "placeholder") if token in full_text]
headings = [p.text for p in doc.paragraphs if p.style and p.style.name.startswith("Heading")]
assert path.stat().st_size > 10_000, "DOCX 文件异常偏小"
assert not missing_features, f"缺少功能编号：{missing_features}"
assert not placeholders, f"存在占位内容：{placeholders}"
assert len(doc.tables) >= 8, "关键分析表格数量不足"
assert any("最终建议" in h for h in headings), "缺少最终建议章节"
assert "月度有效活跃行程数" in full_text, "缺少北极星指标"
print(f"文件：{path}")
print(f"大小：{path.stat().st_size:,} 字节")
print(f"正文段落：{len(texts)}；表格：{len(doc.tables)}；标题：{len(headings)}")
print("功能编号 F01—F14：完整")
print("占位内容检查：通过")
print("关键章节与指标检查：通过")
